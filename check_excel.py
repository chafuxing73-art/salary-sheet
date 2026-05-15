from openpyxl import load_workbook

wb = load_workbook("总工资表.xlsx")
ws = wb.active

print("=== 工资表结构分析 ===")
print(f"总行数: {ws.max_row}")
print(f"总列数: {ws.max_column}")

print("\n--- 第1行表头 ---")
for col in range(1, ws.max_column + 1):
    value = ws.cell(row=1, column=col).value
    print(f"列{col}: {repr(value)}")

print("\n--- 第2行表头 ---")
for col in range(1, ws.max_column + 1):
    value = ws.cell(row=2, column=col).value
    print(f"列{col}: {repr(value)}")

print("\n--- 第3行数据（第一条记录）---")
for col in range(1, ws.max_column + 1):
    value = ws.cell(row=3, column=col).value
    print(f"列{col}: {repr(value)}")