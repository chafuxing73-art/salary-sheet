import os
import sys
import json
import time
import uuid
import threading
import queue
import pandas as pd
import requests
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Side
from flask import Flask, render_template_string, request, jsonify, Response, stream_with_context

app = Flask(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

try:
    from config import CORP_ID, CORP_SECRET, AGENT_ID
except ImportError:
    CORP_ID = ""
    CORP_SECRET = ""
    AGENT_ID = 0

USER_MAP_FILE = os.path.join(BASE_DIR, "总部人员信息表.xlsx")

UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
SPLIT_DIR = os.path.join(BASE_DIR, "工资条拆分文件")
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(SPLIT_DIR, exist_ok=True)

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

log_queues = {}
task_status = {}


class WebLogHandler:
    def __init__(self, task_id):
        self.task_id = task_id

    def info(self, msg):
        self._emit("info", msg)

    def warning(self, msg):
        self._emit("warning", msg)

    def error(self, msg):
        self._emit("error", msg)

    def _emit(self, level, msg):
        entry = {"time": datetime.now().strftime("%H:%M:%S"), "level": level, "msg": msg}
        if self.task_id in log_queues:
            log_queues[self.task_id].put(entry)


def get_access_token(corp_id, corp_secret, log):
    url = f"https://qyapi.weixin.qq.com/cgi-bin/gettoken?corpid={corp_id}&corpsecret={corp_secret}"
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        if data.get("errcode") != 0:
            raise Exception(f"获取access_token失败: {data.get('errmsg')}")
        log.info("✅ 获取企业微信access_token成功")
        return data["access_token"]
    except Exception as e:
        log.error(f"❌ 获取access_token失败: {str(e)}")
        raise


def load_user_map_from_excel(excel_path):
    """从本地Excel文件加载姓名:userid映射"""
    wb = load_workbook(excel_path)
    ws = wb.active
    user_map = {}
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            name = str(row[0])
            if '-' in name:
                name = name.split('-')[-1]
            user_map[name] = str(row[1])
    
    return user_map


def upload_media(access_token, file_path, log):
    url = f"https://qyapi.weixin.qq.com/cgi-bin/media/upload?access_token={access_token}&type=file"
    try:
        with open(file_path, "rb") as f:
            files = {"media": (os.path.basename(file_path), f, "application/octet-stream")}
            resp = requests.post(url, files=files, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        if data.get("errcode") != 0:
            raise Exception(f"上传文件失败: {data.get('errmsg')}")
        return data["media_id"]
    except Exception as e:
        log.error(f"❌ 上传文件失败: {str(e)}")
        raise


def send_file_msg(access_token, agent_id, userid, media_id, name, log):
    url = f"https://qyapi.weixin.qq.com/cgi-bin/message/send?access_token={access_token}"
    payload = {
        "touser": userid,
        "msgtype": "file",
        "agentid": agent_id,
        "file": {"media_id": media_id},
        "safe": 0
    }
    try:
        resp = requests.post(url, json=payload, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        if data.get("errcode") != 0:
            raise Exception(f"发送消息失败: {data.get('errmsg')}")
        log.info(f"✅ {name} 工资条文件发送成功")
        return True
    except Exception as e:
        log.error(f"❌ {name} 发送失败: {str(e)}")
        return False


def send_text_msg(access_token, agent_id, userid, row_data, name, log):
    url = f"https://qyapi.weixin.qq.com/cgi-bin/message/send?access_token={access_token}"
    dept = row_data[1] if pd.notna(row_data[1]) else ""
    post = row_data[2] if pd.notna(row_data[2]) else ""
    real_attend = row_data[6] if pd.notna(row_data[6]) else ""
    base_salary = row_data[7] if pd.notna(row_data[7]) else ""
    post_salary = row_data[8] if pd.notna(row_data[8]) else ""
    performance = row_data[9] if pd.notna(row_data[9]) else ""
    overtime = row_data[10] if pd.notna(row_data[10]) else ""
    house_subsidy = row_data[11] if pd.notna(row_data[11]) else ""
    full_attend = row_data[12] if pd.notna(row_data[12]) else ""
    seniority = row_data[20] if pd.notna(row_data[20]) else ""
    should_pay = row_data[21] if pd.notna(row_data[21]) else ""
    social_security = row_data[22] if pd.notna(row_data[22]) else ""
    should_deduct = row_data[23] if pd.notna(row_data[23]) else ""
    real_pay = row_data[24] if pd.notna(row_data[24]) else ""

    month = datetime.now().strftime("%Y年%m月")
    text_content = f"""【{month} 工资条 - {name}】
部门：{dept}
岗位：{post}
实出勤：{real_attend} 天
-------------------
【工资明细】
基本工资：{base_salary} 元
岗位薪资：{post_salary} 元
绩效工资：{performance} 元
加班工资：{overtime} 元
房补：{house_subsidy} 元
全勤奖：{full_attend} 元
工龄补贴：{seniority} 元
-------------------
应发合计：{should_pay} 元
社保扣款：{social_security} 元
应扣合计：{should_deduct} 元
本月实发：{real_pay} 元
-------------------
如有疑问，请联系财务核对~"""

    payload = {
        "touser": userid,
        "msgtype": "text",
        "agentid": agent_id,
        "text": {"content": text_content},
        "safe": 0
    }
    try:
        resp = requests.post(url, json=payload, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        if data.get("errcode") != 0:
            raise Exception(f"发送消息失败: {data.get('errmsg')}")
        log.info(f"✅ {name} 文字工资条发送成功")
        return True
    except Exception as e:
        log.error(f"❌ {name} 发送失败: {str(e)}")
        return False


def split_all_salary(salary_file, save_dir, log):
    wb_original = load_workbook(salary_file)
    ws_original = wb_original["工资表"]

    header_row1 = [cell.value for cell in ws_original[1]]
    header_row2 = [cell.value for cell in ws_original[2]]

    data_rows = []
    for row in ws_original.iter_rows(min_row=3, values_only=True):
        data_rows.append(row)

    name_col_idx = None
    for idx, cell in enumerate(header_row1):
        if cell == "姓名":
            name_col_idx = idx
            break
    if name_col_idx is None:
        for idx, cell in enumerate(header_row2):
            if cell == "姓名":
                name_col_idx = idx
                break

    if name_col_idx is None:
        raise Exception("❌ 未找到姓名列，请检查工资表结构")

    success_count = 0
    for row in data_rows:
        name = row[name_col_idx]
        if pd.isna(name):
            log.warning("⚠️ 跳过空姓名行")
            continue

        wb_new = Workbook()
        ws_new = wb_new.active
        ws_new.title = f"{name}工资条"

        ws_new.append(header_row1)
        ws_new.append(header_row2)
        ws_new.append(row)

        for row_cells in ws_new.iter_rows():
            for cell in row_cells:
                cell.border = thin_border
                cell.alignment = center_alignment

        for col in ws_new.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length * 1.2 + 2, 20)
            ws_new.column_dimensions[column].width = adjusted_width

        month = datetime.now().strftime("%Y%m")
        file_name = f"工资条-{name}-{month}.xlsx"
        file_path = os.path.join(save_dir, file_name)
        wb_new.save(file_path)
        success_count += 1
        log.info(f"📄 生成工资条：{file_name}")

    log.info(f"✅ 工资条拆分完成，共生成 {success_count} 个文件")
    return data_rows, name_col_idx


def run_task(task_id, salary_file, send_mode, max_retries):
    log = WebLogHandler(task_id)
    try:
        task_status[task_id] = "running"

        log.info("🚀 开始拆分工资条...")
        data_rows, name_col_idx = split_all_salary(salary_file, SPLIT_DIR, log)

        log.info("📋 加载本地用户映射...")
        user_map = load_user_map_from_excel(USER_MAP_FILE)
        log.info(f"✅ 从本地文件读取到 {len(user_map)} 个用户映射")

        log.info("🔑 获取企业微信凭证...")
        access_token = get_access_token(CORP_ID, CORP_SECRET, log)

        success_send = 0
        fail_send = 0
        skip_send = 0
        month = datetime.now().strftime("%Y%m")

        log.info(f"📨 开始发送工资条，模式：{send_mode}，共 {len(data_rows)} 人")

        for row in data_rows:
            name = row[name_col_idx]
            if pd.isna(name):
                skip_send += 1
                continue

            if name not in user_map:
                log.warning(f"⚠️ 用户列表中未找到：{name}，跳过发送")
                skip_send += 1
                continue
            userid = user_map[name]

            send_success = False
            if send_mode == "file":
                file_name = f"工资条-{name}-{month}.xlsx"
                file_path = os.path.join(SPLIT_DIR, file_name)
                if not os.path.exists(file_path):
                    log.warning(f"⚠️ 未找到{name}的工资条文件，跳过")
                    skip_send += 1
                    continue

                for i in range(max_retries):
                    try:
                        media_id = upload_media(access_token, file_path, log)
                        send_success = send_file_msg(access_token, AGENT_ID, userid, media_id, name, log)
                        if send_success:
                            break
                    except Exception as e:
                        log.warning(f"第{i+1}次发送失败，重试中：{str(e)}")

            elif send_mode == "text":
                for i in range(max_retries):
                    send_success = send_text_msg(access_token, AGENT_ID, userid, row, name, log)
                    if send_success:
                        break

            if send_success:
                success_send += 1
            else:
                fail_send += 1
                log.error(f"❌ {name} 最终发送失败，已达最大重试次数")

        log.info("=" * 50)
        log.info("📊 工资条发送任务最终结果")
        log.info(f"✅ 发送成功：{success_send} 人")
        log.info(f"❌ 发送失败：{fail_send} 人")
        log.info(f"⚠️ 跳过发送：{skip_send} 人")
        log.info(f"📄 总员工数：{len(data_rows)} 人")
        log.info("=" * 50)

        task_status[task_id] = "done"
        log_queues[task_id].put({
            "time": datetime.now().strftime("%H:%M:%S"),
            "level": "result",
            "msg": json.dumps({"success": success_send, "fail": fail_send, "skip": skip_send, "total": len(data_rows)})
        })

    except Exception as e:
        log.error(f"❌ 任务执行失败：{str(e)}")
        task_status[task_id] = "error"


@app.route("/")
def index():
    return render_template_string(HTML_TEMPLATE)


@app.route("/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"ok": False, "msg": "未选择文件"}), 400
    f = request.files["file"]
    if not f.filename.endswith((".xlsx", ".xls")):
        return jsonify({"ok": False, "msg": "请上传 .xlsx 或 .xls 文件"}), 400
    save_path = os.path.join(UPLOAD_DIR, f"总工资表_{uuid.uuid4().hex[:8]}.xlsx")
    f.save(save_path)
    return jsonify({"ok": True, "path": save_path, "name": f.filename})


@app.route("/run", methods=["POST"])
def run():
    data = request.json or {}
    salary_file = data.get("salary_file", "").strip()
    send_mode = data.get("send_mode", "file")
    max_retries = int(data.get("max_retries", 3))

    if not CORP_ID or not CORP_SECRET or AGENT_ID == 0:
        return jsonify({"ok": False, "msg": "请先在 config.py 中配置企业微信信息"}), 400
    if not salary_file or not os.path.exists(salary_file):
        return jsonify({"ok": False, "msg": "请先上传工资表文件"}), 400

    task_id = uuid.uuid4().hex[:12]
    log_queues[task_id] = queue.Queue()
    task_status[task_id] = "running"

    t = threading.Thread(
        target=run_task,
        args=(task_id, salary_file, send_mode, max_retries),
        daemon=True
    )
    t.start()

    return jsonify({"ok": True, "task_id": task_id})


@app.route("/logs/<task_id>")
def logs(task_id):
    def generate():
        q = log_queues.get(task_id)
        if not q:
            yield f"data: {json.dumps({'time':'--:--:--','level':'error','msg':'无效任务ID'})}\n\n"
            return
        while True:
            try:
                entry = q.get(timeout=1)
                yield f"data: {json.dumps(entry, ensure_ascii=False)}\n\n"
                if entry.get("level") == "result":
                    break
            except queue.Empty:
                if task_status.get(task_id) in ("done", "error"):
                    yield f"data: {json.dumps({'time':datetime.now().strftime('%H:%M:%S'),'level':'system','msg':'任务已结束'})}\n\n"
                    break

    return Response(stream_with_context(generate()), mimetype="text/event-stream")


HTML_TEMPLATE = r"""
<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>工资条发送系统</title>
<link href="https://cdn.bootcdn.net/ajax/libs/twitter-bootstrap/4.6.2/css/bootstrap.min.css" rel="stylesheet">
<style>
body { background: #f0f2f5; }
.navbar { background: #001529 !important; }
.navbar-brand { color: #fff !important; font-weight: 600; }
.card { border: none; box-shadow: 0 1px 4px rgba(0,0,0,.08); border-radius: 6px; }
.card-header { background: #fafafa; font-weight: 600; border-bottom: 1px solid #f0f0f0; }
.log-box { background: #1e1e1e; color: #d4d4d4; font-family: Consolas, 'Courier New', monospace; font-size: 13px;
  height: 420px; overflow-y: auto; padding: 12px 16px; border-radius: 4px; line-height: 1.7; }
.log-box .log-info { color: #9cdcfe; }
.log-box .log-warning { color: #dcdcaa; }
.log-box .log-error { color: #f44747; }
.log-box .log-system { color: #6a9955; }
.log-box .log-result { color: #4ec9b0; font-weight: bold; }
.stat-card { text-align: center; padding: 20px 10px; }
.stat-card .stat-num { font-size: 36px; font-weight: 700; }
.stat-card .stat-label { font-size: 14px; color: #888; margin-top: 4px; }
.stat-success .stat-num { color: #52c41a; }
.stat-fail .stat-num { color: #f5222d; }
.stat-skip .stat-num { color: #faad14; }
.stat-total .stat-num { color: #1890ff; }
#uploadArea { border: 2px dashed #d9d9d9; border-radius: 6px; padding: 30px; text-align: center; cursor: pointer; transition: border-color .3s; }
#uploadArea:hover, #uploadArea.dragover { border-color: #1890ff; }
#uploadArea .upload-icon { font-size: 40px; color: #1890ff; }
.form-control-sm { font-size: 13px; }
.btn-primary { background: #1890ff; border-color: #1890ff; }
.btn-primary:hover { background: #40a9ff; border-color: #40a9ff; }
.btn-primary:disabled { background: #1890ff; border-color: #1890ff; opacity: .65; }
</style>
</head>
<body>

<nav class="navbar navbar-dark mb-4">
  <a class="navbar-brand" href="#">📋 工资条发送系统</a>
  <span class="navbar-text text-light" style="font-size:13px;">企业微信工资条拆分与发送</span>
</nav>

<div class="container-fluid" style="max-width:1200px;">

  <div class="row">
    <div class="col-lg-6">
      <div class="card mb-4">
        <div class="card-header">📁 上传工资表</div>
        <div class="card-body">
          <div id="uploadArea" onclick="document.getElementById('fileInput').click()">
            <div class="upload-icon">📂</div>
            <div style="margin-top:8px;color:#666;">点击上传 总工资表.xlsx</div>
            <div style="font-size:12px;color:#999;margin-top:4px;">支持 .xlsx / .xls 格式</div>
          </div>
          <input type="file" id="fileInput" accept=".xlsx,.xls" style="display:none">
          <div id="fileInfo" class="mt-2" style="display:none;">
            <span class="badge badge-success">✅ 已上传</span>
            <span id="fileName" style="font-size:13px;margin-left:6px;"></span>
          </div>
        </div>
      </div>

      <div class="card mb-4">
        <div class="card-header">🎯 发送设置</div>
        <div class="card-body">
          <div class="form-group">
            <label>发送模式</label>
            <select class="form-control form-control-sm" id="sendMode">
              <option value="file">文件模式（发送Excel文件）</option>
              <option value="text">文字模式（发送文字工资条）</option>
            </select>
          </div>
          <div class="form-group">
            <label>最大重试次数</label>
            <input type="number" class="form-control form-control-sm" id="maxRetries" value="3" min="1" max="10">
          </div>
          <button class="btn btn-primary btn-block mt-3" id="runBtn" onclick="startRun()">
            🚀 一键执行：拆分 + 发送
          </button>
        </div>
      </div>
    </div>

    <div class="col-lg-6">
      <div class="card mb-4">
        <div class="card-header d-flex justify-content-between align-items-center">
          <span>📋 运行日志</span>
          <button class="btn btn-sm btn-outline-secondary" onclick="clearLog()">清空</button>
        </div>
        <div class="card-body p-0">
          <div class="log-box" id="logBox">
            <div class="log-system">等待执行...</div>
          </div>
        </div>
      </div>

      <div class="card mb-4" id="resultCard" style="display:none;">
        <div class="card-header">📊 发送结果统计</div>
        <div class="card-body">
          <div class="row">
            <div class="col-3 stat-card stat-success">
              <div class="stat-num" id="statSuccess">0</div>
              <div class="stat-label">发送成功</div>
            </div>
            <div class="col-3 stat-card stat-fail">
              <div class="stat-num" id="statFail">0</div>
              <div class="stat-label">发送失败</div>
            </div>
            <div class="col-3 stat-card stat-skip">
              <div class="stat-num" id="statSkip">0</div>
              <div class="stat-label">跳过发送</div>
            </div>
            <div class="col-3 stat-card stat-total">
              <div class="stat-num" id="statTotal">0</div>
              <div class="stat-label">总员工数</div>
            </div>
          </div>
        </div>
      </div>
    </div>
  </div>
</div>

<script src="https://cdn.bootcdn.net/ajax/libs/jquery/3.6.4/jquery.min.js"></script>
<script>
var uploadedFilePath = "";

$("#fileInput").on("change", function() {
  var file = this.files[0];
  if (!file) return;
  var fd = new FormData();
  fd.append("file", file);
  $.ajax({
    url: "/upload",
    type: "POST",
    data: fd,
    processData: false,
    contentType: false,
    success: function(res) {
      if (res.ok) {
        uploadedFilePath = res.path;
        $("#fileName").text(res.name);
        $("#fileInfo").show();
      } else {
        alert(res.msg || "上传失败");
      }
    },
    error: function() { alert("上传请求失败"); }
  });
});

var uploadArea = document.getElementById("uploadArea");
uploadArea.addEventListener("dragover", function(e) { e.preventDefault(); this.classList.add("dragover"); });
uploadArea.addEventListener("dragleave", function() { this.classList.remove("dragover"); });
uploadArea.addEventListener("drop", function(e) {
  e.preventDefault();
  this.classList.remove("dragover");
  var files = e.dataTransfer.files;
  if (files.length > 0) {
    $("#fileInput")[0].files = files;
    $("#fileInput").trigger("change");
  }
});

function appendLog(time, level, msg) {
  var box = $("#logBox");
  var cls = "log-" + level;
  box.append('<div class="' + cls + '">[' + time + '] ' + escapeHtml(msg) + '</div>');
  box.scrollTop(box[0].scrollHeight);
}

function escapeHtml(text) {
  var d = document.createElement("div");
  d.textContent = text;
  return d.innerHTML;
}

function clearLog() {
  $("#logBox").html('<div class="log-system">日志已清空</div>');
  $("#resultCard").hide();
}

function startRun() {
  var sendMode = $("#sendMode").val();
  var maxRetries = $("#maxRetries").val();

  if (!uploadedFilePath) {
    alert("请先上传工资表文件");
    return;
  }

  $("#runBtn").prop("disabled", true).text("⏳ 执行中...");
  $("#resultCard").hide();
  clearLog();
  appendLog("--:--:--", "system", "🚀 任务启动中...");

  $.ajax({
    url: "/run",
    type: "POST",
    contentType: "application/json",
    data: JSON.stringify({
      salary_file: uploadedFilePath,
      send_mode: sendMode,
      max_retries: parseInt(maxRetries)
    }),
    success: function(res) {
      if (res.ok) {
        listenLogs(res.task_id);
      } else {
        appendLog("--:--:--", "error", "启动失败：" + (res.msg || "未知错误"));
        $("#runBtn").prop("disabled", false).text("🚀 一键执行：拆分 + 发送");
      }
    },
    error: function() {
      appendLog("--:--:--", "error", "请求失败，请检查网络");
      $("#runBtn").prop("disabled", false).text("🚀 一键执行：拆分 + 发送");
    }
  });
}

function listenLogs(taskId) {
  var source = new EventSource("/logs/" + taskId);
  source.onmessage = function(e) {
    var entry = JSON.parse(e.data);
    if (entry.level === "result") {
      var r = JSON.parse(entry.msg);
      $("#statSuccess").text(r.success);
      $("#statFail").text(r.fail);
      $("#statSkip").text(r.skip);
      $("#statTotal").text(r.total);
      $("#resultCard").show();
      appendLog(entry.time, "result", "📊 任务完成！成功:" + r.success + " 失败:" + r.fail + " 跳过:" + r.skip + " 总计:" + r.total);
      source.close();
      $("#runBtn").prop("disabled", false).text("🚀 一键执行：拆分 + 发送");
    } else {
      appendLog(entry.time, entry.level, entry.msg);
    }
  };
  source.onerror = function() {
    source.close();
    $("#runBtn").prop("disabled", false).text("🚀 一键执行：拆分 + 发送");
  };
}
</script>
</body>
</html>
"""


if __name__ == "__main__":
    print("=" * 50)
    print("📋 工资条发送系统已启动")
    print("🌐 本机访问: http://127.0.0.1:5000")
    print("🌐 局域网访问: http://<本机IP>:5000")
    print("=" * 50)
    app.run(host="0.0.0.0", port=5000, debug=False)
